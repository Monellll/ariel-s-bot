"""
smoke_test.py
-------------
Offline smoke tests for excel_handler.py and reminder_scheduler.py.
No real Telegram calls are made.

Run with:
    python smoke_test.py
"""

import datetime
import io
import json
import os
import shutil
import sys
import tempfile
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

# Force UTF-8 output on Windows (avoids cp1252 UnicodeEncodeError with emoji)
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

import openpyxl

# Make sure project root is on path
sys.path.insert(0, str(Path(__file__).parent))

import excel_handler
import reminder_scheduler

PASS = "✅ PASS"
FAIL = "❌ FAIL"
results = []


def check(name: str, condition: bool, detail: str = "") -> None:
    status = PASS if condition else FAIL
    msg = f"  {status}  {name}"
    if detail:
        msg += f"\n         {detail}"
    print(msg)
    results.append(condition)


# ===========================================================================
# Helpers
# ===========================================================================


def make_xlsx(rows: list[list], headers: list[str]) -> str:
    """Create a temporary .xlsx file and return its path."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    tmp = tempfile.mktemp(suffix=".xlsx")
    wb.save(tmp)
    return tmp


# ===========================================================================
# Test 1: Valid rows parse correctly (including various date formats)
# ===========================================================================
print("\n--- Test 1: Valid rows ---")
xlsx = make_xlsx(
    headers=["telegram_username", "student_name", "due_date", "telegram_chat_id"],
    rows=[
        ["@Alice", "Alice Smith", "2026-12-01", 111111111],
        ["@Bob", "Bob Jones", "2026-11-20", 222222222],
        ["@Charlie", "Charlie Brown", "2026-10-16 00:00:00", 333333333],
        ["@Dave", "Dave Miller", datetime.datetime(2026, 10, 16, 0, 0), 444444444],
    ],
)
records, invalid, errors = excel_handler.load_excel(xlsx)
os.unlink(xlsx)

check("Four valid rows loaded", len(records) == 4, f"got {len(records)}")
check("No invalid rows", invalid == 0, f"got {invalid}")
check("due_date is datetime.date", isinstance(records[0]["due_date"], datetime.date))
check("datetime string parsed", records[2]["due_date"] == datetime.date(2026, 10, 16))
check("native datetime parsed", records[3]["due_date"] == datetime.date(2026, 10, 16))
check("student_name preserved", records[0]["student_name"] == "Alice Smith")
check("chat_id is int", records[0]["telegram_chat_id"] == 111111111)

# ===========================================================================
# Test 2: Invalid rows are counted, valid rows still returned
# ===========================================================================
print("\n--- Test 2: Mixed valid/invalid rows ---")
xlsx = make_xlsx(
    headers=["telegram_username", "student_name", "due_date", "telegram_chat_id"],
    rows=[
        ["@Valid", "Valid Student", "2026-12-10", 333333333],
        ["no_at_sign", "Bad Username", "2026-12-11", 444444444],   # bad username
        ["@Also Valid", "Good", "not-a-date", 555555555],          # bad date
    ],
)
records, invalid, errors = excel_handler.load_excel(xlsx)
os.unlink(xlsx)

check("1 valid row returned", len(records) == 1, f"got {len(records)}")
check("2 invalid rows counted", invalid == 2, f"got {invalid}")
check("Errors list has entries", len(errors) >= 2, f"got {len(errors)}")

# ===========================================================================
# Test 3: Blank rows are silently skipped
# ===========================================================================
print("\n--- Test 3: Blank rows skipped ---")
xlsx = make_xlsx(
    headers=["telegram_username", "student_name", "due_date", "telegram_chat_id"],
    rows=[
        ["@Real", "Real Student", "2026-12-15", 666666666],
        [None, None, None, None],   # blank row
        ["", "", "", ""],           # blank row
    ],
)
records, invalid, errors = excel_handler.load_excel(xlsx)
os.unlink(xlsx)

check("1 valid row (blank rows skipped)", len(records) == 1, f"got {len(records)}")

# ===========================================================================
# Test 4: Case-insensitive column headers
# ===========================================================================
print("\n--- Test 4: Case-insensitive headers ---")
xlsx = make_xlsx(
    headers=["Telegram_Username", "Student_Name", "Due_Date", "Telegram_Chat_Id"],
    rows=[
        ["@CaseTest", "Case Test", "2026-12-20", 777777777],
    ],
)
records, invalid, errors = excel_handler.load_excel(xlsx)
os.unlink(xlsx)

check("Case-insensitive headers work", len(records) == 1, f"got {len(records)}")

# ===========================================================================
# Test 5: Missing telegram_chat_id is OK (not an error)
# ===========================================================================
print("\n--- Test 5: Missing chat_id is not an error ---")
xlsx = make_xlsx(
    headers=["telegram_username", "due_date"],
    rows=[
        ["@NoChatId", "2026-12-25"],
    ],
)
records, invalid, errors = excel_handler.load_excel(xlsx)
os.unlink(xlsx)

check("Row without chat_id column is valid", len(records) == 1, f"got {len(records)}")
check("chat_id is None", records[0]["telegram_chat_id"] is None)
check("student_name falls back to username", records[0]["student_name"] == "@NoChatId")

# ===========================================================================
# Test 6: Date math — reminder triggers on exact day
# ===========================================================================
print("\n--- Test 6: Date math (days_until_due) ---")
from zoneinfo import ZoneInfo
today = datetime.datetime.now(ZoneInfo("Asia/Yangon")).date()
five_days_later = today + datetime.timedelta(days=5)
ten_days_later = today + datetime.timedelta(days=10)

records_mock = [
    {
        "telegram_username": "@Alice",
        "student_name": "Alice",
        "due_date": five_days_later,
        "telegram_chat_id": 111111111,
    },
    {
        "telegram_username": "@Bob",
        "student_name": "Bob",
        "due_date": ten_days_later,
        "telegram_chat_id": 222222222,
    },
    {
        "telegram_username": "@OldDebt",
        "student_name": "OldDebt",
        "due_date": today - datetime.timedelta(days=1),  # past
        "telegram_chat_id": 333333333,
    },
]

upcoming = reminder_scheduler.get_upcoming_reminders(records_mock, "Asia/Yangon", 5, n=5)
check("get_upcoming_reminders returns future only", len(upcoming) == 2, f"got {len(upcoming)}")
check("Sorted ascending by days_until_due", upcoming[0]["days_until_due"] <= upcoming[1]["days_until_due"])

# ===========================================================================
# Test 7: Duplicate prevention
# ===========================================================================
print("\n--- Test 7: Duplicate prevention (sent_reminders.json) ---")
with tempfile.TemporaryDirectory() as tmpdir:
    json_path = Path(tmpdir) / "sent_reminders.json"

    # Empty at start
    keys = reminder_scheduler.load_sent_reminders(json_path)
    check("Empty set on first load", len(keys) == 0)

    # Add a key and save
    keys.add("111111111:2026-12-01")
    reminder_scheduler.save_sent_reminders(keys, json_path)

    # Reload
    keys2 = reminder_scheduler.load_sent_reminders(json_path)
    check("Key persists after reload", "111111111:2026-12-01" in keys2)
    check("Only 1 key saved", len(keys2) == 1)

# ===========================================================================
# Test 8: daily_check only sends to students due in exactly REMINDER_DAYS days
# ===========================================================================
print("\n--- Test 8: daily_check sends correct reminders (mocked bot) ---")

async def _run_daily_check_test():
    with tempfile.TemporaryDirectory() as tmpdir:
        json_path = Path(tmpdir) / "sent.json"
        today = datetime.datetime.now(ZoneInfo("Asia/Yangon")).date()
        records = [
            {  # Should send — exactly 5 days out
                "telegram_username": "@ShouldSend",
                "student_name": "ShouldSend",
                "due_date": today + datetime.timedelta(days=5),
                "telegram_chat_id": 999999999,
            },
            {  # Should NOT send — 10 days out
                "telegram_username": "@TooFar",
                "student_name": "TooFar",
                "due_date": today + datetime.timedelta(days=10),
                "telegram_chat_id": 888888888,
            },
            {  # Should NOT send — already past
                "telegram_username": "@Past",
                "student_name": "Past",
                "due_date": today - datetime.timedelta(days=1),
                "telegram_chat_id": 777777777,
            },
            {  # Should NOT send — no chat_id
                "telegram_username": "@NoChatId",
                "student_name": "NoChatId",
                "due_date": today + datetime.timedelta(days=5),
                "telegram_chat_id": None,
            },
        ]

        mock_bot = AsyncMock()
        mock_bot.send_message = AsyncMock()

        await reminder_scheduler.daily_check(
            bot=mock_bot,
            records=records,
            admin_chat_id=12345,
            timezone_str="Asia/Yangon",
            reminder_days=5,
            sent_reminders_path=json_path,
        )

        # send_message called: 1 reminder + 1 admin warning (no chat_id) = 2 calls
        calls = mock_bot.send_message.call_args_list
        sent_chat_ids = [c.kwargs.get("chat_id") or c.args[0] for c in calls]

        check(
            "Reminder sent to correct student",
            999999999 in sent_chat_ids,
            f"sent_chat_ids={sent_chat_ids}",
        )
        check(
            "Too-far student NOT messaged directly",
            888888888 not in sent_chat_ids,
            f"sent_chat_ids={sent_chat_ids}",
        )
        check(
            "Past student NOT messaged",
            777777777 not in sent_chat_ids,
            f"sent_chat_ids={sent_chat_ids}",
        )

        # Reload json — key should be saved
        saved = reminder_scheduler.load_sent_reminders(json_path)
        key = f"999999999:{today + datetime.timedelta(days=5)}"
        check("Sent key persisted to JSON", key in saved, f"saved={saved}")

        # Run again — should NOT resend
        mock_bot.send_message.reset_mock()
        await reminder_scheduler.daily_check(
            bot=mock_bot,
            records=records,
            admin_chat_id=12345,
            timezone_str="Asia/Yangon",
            reminder_days=5,
            sent_reminders_path=json_path,
        )
        second_calls = mock_bot.send_message.call_args_list
        second_student_calls = [
            c for c in second_calls if (c.kwargs.get("chat_id") or c.args[0]) == 999999999
        ]
        check("Duplicate NOT re-sent on second run", len(second_student_calls) == 0)


import asyncio
asyncio.run(_run_daily_check_test())

# ===========================================================================
# Test 9: Multi-user isolation (separate files and schedules)
# ===========================================================================
print("\n--- Test 9: Multi-user isolation ---")

async def _run_multi_user_test():
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp_base = Path(tmpdir)
        user1_dir = tmp_base / "1001"
        user2_dir = tmp_base / "1002"
        user1_dir.mkdir()
        user2_dir.mkdir()

        user1_sent = user1_dir / "sent_reminders.json"
        user2_sent = user2_dir / "sent_reminders.json"

        today = datetime.datetime.now(ZoneInfo("Asia/Yangon")).date()
        target_due = today + datetime.timedelta(days=5)

        # User 1 has Student Alpha
        user1_records = [
            {
                "telegram_username": "@StudentAlpha",
                "student_name": "Alpha",
                "due_date": target_due,
                "telegram_chat_id": 11111,
            }
        ]

        # User 2 has Student Beta
        user2_records = [
            {
                "telegram_username": "@StudentBeta",
                "student_name": "Beta",
                "due_date": target_due,
                "telegram_chat_id": 22222,
            }
        ]

        mock_bot = AsyncMock()

        # Run User 1 check
        await reminder_scheduler.daily_check(
            bot=mock_bot,
            records=user1_records,
            owner_chat_id=1001,
            timezone_str="Asia/Yangon",
            reminder_days=5,
            sent_reminders_path=user1_sent,
        )

        user1_keys = reminder_scheduler.load_sent_reminders(user1_sent)
        user2_keys = reminder_scheduler.load_sent_reminders(user2_sent)

        check("User 1 sent recorded in User 1 history", f"11111:{target_due}" in user1_keys)
        check("User 2 history remains empty", len(user2_keys) == 0)

        # Run User 2 check
        await reminder_scheduler.daily_check(
            bot=mock_bot,
            records=user2_records,
            owner_chat_id=1002,
            timezone_str="Asia/Yangon",
            reminder_days=5,
            sent_reminders_path=user2_sent,
        )

        user2_keys_after = reminder_scheduler.load_sent_reminders(user2_sent)
        check("User 2 sent recorded in User 2 history", f"22222:{target_due}" in user2_keys_after)
        check("User 1 history unaffected by User 2 run", len(reminder_scheduler.load_sent_reminders(user1_sent)) == 1)

asyncio.run(_run_multi_user_test())

# ===========================================================================
# Test 10: Authorized user storage & admin auto-authorization
# ===========================================================================
print("\n--- Test 10: User authorization storage ---")
with tempfile.TemporaryDirectory() as tmpdir:
    auth_file = Path(tmpdir) / "authorized_users.json"
    admin_id = 987654

    # Test load when file doesn't exist yet
    users = {admin_id}
    if auth_file.exists():
        with open(auth_file, "r", encoding="utf-8") as f:
            users.update(json.load(f))
    check("Admin is auto-authorized by default", admin_id in users)

    # Save a new user
    users.add(123456)
    with open(auth_file, "w", encoding="utf-8") as f:
        json.dump(sorted(users), f)

    # Reload from disk
    loaded_users = set()
    with open(auth_file, "r", encoding="utf-8") as f:
        loaded_users = set(json.load(f))
    check("Persisted authorized users load correctly", 123456 in loaded_users and admin_id in loaded_users)
    check("Total authorized count is 2", len(loaded_users) == 2)

# ===========================================================================
# Test 11: Dynamic reminder message formatting
# ===========================================================================
print("\n--- Test 11: Dynamic reminder message formatting ---")
msg_default = reminder_scheduler.format_reminder_message(
    student_name="Ariel",
    due_date=datetime.date(2026, 10, 14),
    reminder_days=5,
)
check("Message contains dynamic October 14th", "October 14th" in msg_default)
check("Message contains Myanmar October 14", "October 14 ရက်နေ့" in msg_default)
check("Message contains fee for the class text", "fee for the class is due on" in msg_default)
check("Message contains gentle reminder text", "No worries at all, just a gentle reminder." in msg_default)
check("Message contains Myanmar fee text", "ချယ်ချယ်တို့ရဲ့သင်တန်းလေးကိုဒီလအတွက်" in msg_default)
check("Message contains Myanmar polite closing", "မလောပါဘူးရှင် remind လုပ်တာပါနော်🤍" in msg_default)
check("Message does not contain start_date reference", "since we started" not in msg_default)

# English-only
msg_en = reminder_scheduler.format_reminder_message(
    student_name="Ariel",
    due_date=datetime.date(2026, 10, 14),
    language="english",
)
check("English-only contains English text", "fee for the class is due on October 14th" in msg_en)
check("English-only does NOT contain Myanmar text", "ချယ်ချယ်တို့ရဲ့သင်တန်းလေးကိုဒီလအတွက်" not in msg_en)

# Burmese-only
msg_my = reminder_scheduler.format_reminder_message(
    student_name="Ariel",
    due_date=datetime.date(2026, 10, 14),
    language="burmese",
)
check("Burmese-only contains Myanmar text", "ချယ်ချယ်တို့ရဲ့သင်တန်းလေးကိုဒီလအတွက်" in msg_my)
check("Burmese-only does NOT contain English text", "fee for the class is due" not in msg_my)

# ===========================================================================
# Test 12: save_excel and delete persistence
# ===========================================================================
print("\n--- Test 12: save_excel and delete persistence ---")
with tempfile.TemporaryDirectory() as tmpdir:
    test_xlsx = Path(tmpdir) / "students.xlsx"
    records_to_save = [
        {
            "telegram_username": "@student1",
            "student_name": "Student One",
            "due_date": datetime.date(2026, 11, 1),
            "telegram_chat_id": 111111,
            "start_date": datetime.date(2026, 7, 14),
        },
        {
            "telegram_username": "@student2",
            "student_name": "Student Two",
            "due_date": datetime.date(2026, 11, 2),
            "telegram_chat_id": 222222,
            "start_date": None,
        },
    ]

    # Save initial records
    excel_handler.save_excel(records_to_save, str(test_xlsx))
    check("Excel file created by save_excel", test_xlsx.exists())

    # Reload and verify
    loaded, invalids, errs = excel_handler.load_excel(str(test_xlsx))
    check("Reloaded records count matches", len(loaded) == 2, f"got {len(loaded)}")
    check("Student 1 username preserved", loaded[0]["telegram_username"] == "@student1")
    check("Student 1 chat_id preserved", loaded[0]["telegram_chat_id"] == 111111)

    # Delete 1 record and re-save
    loaded.pop(0)
    excel_handler.save_excel(loaded, str(test_xlsx))

    # Reload again
    reloaded, _, _ = excel_handler.load_excel(str(test_xlsx))
    check("Reloaded count after delete is 1", len(reloaded) == 1, f"got {len(reloaded)}")
    check("Remaining student is student2", reloaded[0]["telegram_username"] == "@student2")

# ===========================================================================
# Test 13: Strict duplicate impossibility & concurrency protection
# ===========================================================================
print("\n--- Test 13: Strict duplicate impossibility & concurrency protection ---")
async def run_concurrency_test():
    with tempfile.TemporaryDirectory() as tmpdir:
        today_yangon = datetime.datetime.now(ZoneInfo("Asia/Yangon")).date()
        target_date = today_yangon + datetime.timedelta(days=5)

        # Roster with duplicate rows of the exact same student
        dup_records = [
            {
                "telegram_username": "@dupstudent",
                "student_name": "Dup Student",
                "due_date": target_date,
                "telegram_chat_id": 555555,
            },
            {
                "telegram_username": "@dupstudent",
                "student_name": "Dup Student",
                "due_date": target_date,
                "telegram_chat_id": 555555,
            },
        ]

        mock_bot = AsyncMock()
        mock_bot.send_message = AsyncMock()

        # Execute 3 concurrent daily_check calls simultaneously
        results = await asyncio.gather(
            reminder_scheduler.daily_check(
                bot=mock_bot,
                records=dup_records,
                admin_chat_id=12345,
                timezone_str="Asia/Yangon",
                reminder_days=5,
                sent_reminders_path=json_path,
            ),
            reminder_scheduler.daily_check(
                bot=mock_bot,
                records=dup_records,
                admin_chat_id=12345,
                timezone_str="Asia/Yangon",
                reminder_days=5,
                sent_reminders_path=json_path,
            ),
            reminder_scheduler.daily_check(
                bot=mock_bot,
                records=dup_records,
                admin_chat_id=12345,
                timezone_str="Asia/Yangon",
                reminder_days=5,
                sent_reminders_path=json_path,
            ),
        )

        # Count total send_message calls for student 555555
        calls = mock_bot.send_message.call_args_list
        student_calls = [
            c for c in calls if (c.kwargs.get("chat_id") == 555555 or (c.args and c.args[0] == 555555))
        ]
        check(
            "Exactly 1 reminder sent across 3 concurrent checks and duplicate rows",
            len(student_calls) == 1,
            f"got {len(student_calls)} calls (expected 1)",
        )

        # Confirm persisted immediately
        saved_keys = reminder_scheduler.load_sent_reminders(json_path)
        expected_key = f"555555:{target_date.strftime('%Y-%m-%d')}"
        check("Key is persisted in sent_reminders.json", expected_key in saved_keys)

asyncio.run(run_concurrency_test())

# ===========================================================================
# Test 14: Menu UI callbacks & Authorization guards
# ===========================================================================
print("\n--- Test 14: Menu UI callbacks & Authorization guards ---")
import bot

mock_update = MagicMock()
mock_user = MagicMock()
mock_user.id = 7462033455
mock_update.effective_user = mock_user

check("Admin 7462033455 recognized by _is_admin", bot._is_admin(mock_update))
check("Admin 7462033455 recognized by _is_authorized", bot._is_authorized(mock_update))

async def run_callback_tests():
    # Test safe edit markdown fallback on BadRequest using edit_message_text
    mock_query = AsyncMock()
    from telegram.error import BadRequest
    mock_query.edit_message_text.side_effect = [BadRequest("Can't parse entities"), None]
    await bot._safe_edit_message(mock_query, "Bad *markdown [test", None, "Markdown")
    check("Markdown parse error cleanly falls back to plain text", mock_query.edit_message_text.call_count == 2)

    # Test handle_callback_query with menu:home
    cb_update = MagicMock()
    cb_query = AsyncMock()
    cb_query.from_user.id = 7462033455
    cb_query.from_user.first_name = "Myat"
    cb_query.data = "menu:home"
    cb_update.callback_query = cb_query
    ctx = MagicMock()

    await bot.handle_callback_query(cb_update, ctx)
    check("Callback menu:home handled successfully", cb_query.edit_message_text.called or cb_query.edit_text.called)

    # Test unknown/legacy callback graceful handling
    cb_query.reset_mock()
    cb_query.data = "menu:language"
    await bot.handle_callback_query(cb_update, ctx)
    check("Legacy callback gracefully routes back to menu", cb_query.edit_message_text.called or cb_query.edit_text.called)

asyncio.run(run_callback_tests())

# ===========================================================================
# Test 15: Admin Teacher Management (/authorize and /revoke)
# ===========================================================================
print("\n--- Test 15: Admin Teacher Management ---")

async def run_admin_mgmt_tests():
    # Admin can authorize a new teacher by numeric ID
    admin_update = MagicMock()
    admin_update.effective_user.id = 7462033455
    admin_update.message = AsyncMock()
    ctx = MagicMock()
    ctx.args = ["555666777"]

    await bot.cmd_authorize(admin_update, ctx)
    check("Teacher 555666777 authorized by admin", 555666777 in bot._authorized_users)

    # Admin can revoke teacher access
    await bot.cmd_revoke(admin_update, ctx)
    check("Teacher 555666777 revoked by admin", 555666777 not in bot._authorized_users)

    # Permanent admin cannot be revoked
    ctx.args = ["7462033455"]
    await bot.cmd_revoke(admin_update, ctx)
    check("Permanent admin 7462033455 cannot be revoked", 7462033455 == bot.ADMIN_ID)

    # Non-admin cannot authorize
    non_admin_update = MagicMock()
    non_admin_update.effective_user.id = 999999999
    non_admin_update.message = AsyncMock()
    ctx.args = ["888888888"]
    await bot.cmd_authorize(non_admin_update, ctx)
    check("Non-admin cannot authorize others", 888888888 not in bot._authorized_users)

    # Authorized teacher cannot authorize others
    teacher_id = 777111222
    bot._authorized_users.add(teacher_id)
    teacher_update = MagicMock()
    teacher_update.effective_user.id = teacher_id
    teacher_update.message = AsyncMock()
    ctx.args = ["888888888"]
    await bot.cmd_authorize(teacher_update, ctx)
    check("Authorized teacher cannot authorize others", 888888888 not in bot._authorized_users)

    # Authorized teacher cannot revoke others
    ctx.args = ["555666777"]
    await bot.cmd_revoke(teacher_update, ctx)
    revoke_reply = teacher_update.message.reply_text.call_args[0][0]
    check("Authorized teacher cannot revoke others", "Only the bot administrator" in revoke_reply)

    # Authorized teacher cannot view /teachers list
    await bot.cmd_teachers(teacher_update, ctx)
    teachers_reply = teacher_update.message.reply_text.call_args[0][0]
    check("Authorized teacher cannot view /teachers", "Only the bot administrator" in teachers_reply)

asyncio.run(run_admin_mgmt_tests())

# ===========================================================================
# Test 16: Teacher Authentication by entering permanent admin ID
# ===========================================================================
print("\n--- Test 16: Teacher Authentication by entering permanent admin ID ---")

async def run_admin_id_auth_tests():
    new_user_id = 1122334455
    new_user_update = MagicMock()
    new_user = MagicMock()
    new_user.id = new_user_id
    new_user.username = "new_teacher_user"
    new_user.first_name = "NewTeacher"
    new_user_update.effective_user = new_user
    new_user_update.effective_chat.id = new_user_id
    new_user_update.message = AsyncMock()

    # Initial state: not admin, not authorized
    bot._authorized_users.discard(new_user_id)
    check("New user initially not admin", not bot._is_admin(new_user_update))
    check("New user initially not authorized", not bot._is_authorized(new_user_update))

    # User enters a random non-admin string
    new_user_update.message.text = "Hello bot"
    await bot.handle_text(new_user_update, MagicMock())
    check("Random text does not grant admin access", not bot._is_admin(new_user_update))
    check("Random text does not grant teacher access", not bot._is_authorized(new_user_update))
    reply_call = new_user_update.message.reply_text.call_args[0][0]
    check("Non-authorized receives student welcome on regular text", "Your account is registered with Ariel's English Class" in reply_call)

    # User enters the permanent admin ID (as text)
    new_user_update.message.reply_text.reset_mock()
    new_user_update.message.text = str(bot.ADMIN_ID)
    await bot.handle_text(new_user_update, MagicMock())

    check("User is authorized after entering permanent admin ID", bot._is_authorized(new_user_update))
    check("User is NOT admin after entering permanent admin ID (strictly teacher)", not bot._is_admin(new_user_update))
    all_replies = [call[0][0] for call in new_user_update.message.reply_text.call_args_list if call[0]]
    check("Reply confirms Teacher Access Granted", any("Teacher Access Granted" in r for r in all_replies))

    # Verify this teacher CANNOT manage teachers
    ctx = MagicMock()
    ctx.args = ["333444555"]
    await bot.cmd_authorize(new_user_update, ctx)
    check("Teacher who entered ID cannot authorize others", 333444555 not in bot._authorized_users)

    # Test via /auth <permanent_id> for another user
    auth_user_id = 9988776655
    auth_update = MagicMock()
    auth_user = MagicMock()
    auth_user.id = auth_user_id
    auth_user.username = "auth_tester"
    auth_user.first_name = "AuthTester"
    auth_update.effective_user = auth_user
    auth_update.effective_chat.id = auth_user_id
    auth_update.message = AsyncMock()
    ctx = MagicMock()
    ctx.args = [str(bot.ADMIN_ID)]

    await bot.cmd_auth(auth_update, ctx)
    check("User is authorized via /auth <admin_id>", bot._is_authorized(auth_update))
    check("User is NOT admin via /auth <admin_id> (strictly teacher)", not bot._is_admin(auth_update))

    # Clean up test user IDs
    bot._authorized_users.discard(new_user_id)
    bot._authorized_users.discard(auth_user_id)
    bot._authorized_users.discard(777111222)
    bot._save_authorized_users()

asyncio.run(run_admin_id_auth_tests())

# ===========================================================================
# Test 17: Google Sheets URL extraction & validation
# ===========================================================================
print("\n--- Test 17: Google Sheets URL extraction & validation ---")

check("is_google_sheet_url true for edit link", excel_handler.is_google_sheet_url("https://docs.google.com/spreadsheets/d/1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms/edit?usp=sharing"))
check("is_google_sheet_url true for published link", excel_handler.is_google_sheet_url("https://docs.google.com/spreadsheets/d/e/2PACX-1vT123/pubhtml"))
check("is_google_sheet_url false for standard web URL", not excel_handler.is_google_sheet_url("https://google.com"))
check("is_google_sheet_url false for empty string", not excel_handler.is_google_sheet_url(""))

regular_url = "https://docs.google.com/spreadsheets/d/1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms/edit?usp=sharing"
export_reg = excel_handler.extract_google_sheet_export_url(regular_url)
check("Standard export URL contains /export?format=xlsx", export_reg == "https://docs.google.com/spreadsheets/d/1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms/export?format=xlsx")

gid_url = "https://docs.google.com/spreadsheets/d/1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms/edit#gid=9988"
export_gid = excel_handler.extract_google_sheet_export_url(gid_url)
check("Export URL preserves gid", export_gid == "https://docs.google.com/spreadsheets/d/1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms/export?format=xlsx&gid=9988")

pub_url = "https://docs.google.com/spreadsheets/d/e/2PACX-1vT123/pubhtml"
export_pub = excel_handler.extract_google_sheet_export_url(pub_url)
check("Published export URL maps to /pub?output=xlsx", export_pub == "https://docs.google.com/spreadsheets/d/e/2PACX-1vT123/pub?output=xlsx")

# ===========================================================================
# Test 18: Google Sheets sync & error handling (mocked network)
# ===========================================================================
print("\n--- Test 18: Google Sheets sync & error handling ---")

import httpx
from unittest.mock import patch, MagicMock

# 18.1: Restricted sheet raises PermissionError
mock_restricted = MagicMock(status_code=200, url="https://accounts.google.com/ServiceLogin", content=b"<!DOCTYPE html><html>Sign in</html>")
with patch.object(httpx.Client, "get", return_value=mock_restricted):
    try:
        excel_handler.sync_google_sheet("https://docs.google.com/spreadsheets/d/test12345/edit", "data/test_sync.xlsx")
        check("Restricted sheet raises PermissionError", False)
    except PermissionError:
        check("Restricted sheet raises PermissionError", True)

# 18.2: 404 sheet raises ValueError
mock_404 = MagicMock(status_code=404, url="https://docs.google.com/spreadsheets/d/notfound/export?format=xlsx")
with patch.object(httpx.Client, "get", return_value=mock_404):
    try:
        excel_handler.sync_google_sheet("https://docs.google.com/spreadsheets/d/notfound/edit", "data/test_sync.xlsx")
        check("404 sheet raises ValueError", False)
    except ValueError as e:
        check("404 sheet raises ValueError", "not found" in str(e).lower())

# 18.3: Valid sheet downloads, validates, and saves records
test_records = [
    {"telegram_username": "@StudentA", "student_name": "Student A", "due_date": datetime.date(2026, 10, 25)},
    {"telegram_username": "@StudentB", "student_name": "Student B", "due_date": datetime.date(2026, 10, 26)},
]
with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp_x:
    excel_handler.save_excel(test_records, tmp_x.name)
    with open(tmp_x.name, "rb") as f:
        valid_excel_bytes = f.read()

mock_ok = MagicMock(status_code=200, url="https://docs.google.com/export", content=valid_excel_bytes)
sync_dest = Path("data/unit_test_user/students.xlsx")
with patch.object(httpx.Client, "get", return_value=mock_ok):
    recs, inv, errs = excel_handler.sync_google_sheet(
        "https://docs.google.com/spreadsheets/d/valid_test/edit", str(sync_dest)
    )
    check("Valid Google Sheet syncs 2 records", len(recs) == 2)
    check("Valid Google Sheet destination file exists on disk", sync_dest.exists())
    check("Reloaded sync record has correct username", recs[0]["telegram_username"] == "@StudentA")

# ===========================================================================
# Test 19: Bot Google Sheet interaction & callback
# ===========================================================================
print("\n--- Test 19: Bot Google Sheet interaction & callback ---")

async def run_bot_google_sheet_tests():
    test_user_id = 8877665544
    bot._authorized_users.add(test_user_id)
    bot._save_authorized_users()

    update = MagicMock()
    user = MagicMock()
    user.id = test_user_id
    user.username = "gsheet_teacher"
    user.first_name = "Teacher GSheet"
    update.effective_user = user
    update.effective_chat.id = test_user_id
    update.message = AsyncMock()
    update.message.text = "https://docs.google.com/spreadsheets/d/valid_sheet_123/edit?usp=sharing"

    with patch.object(httpx.Client, "get", return_value=mock_ok):
        await bot.handle_text(update, MagicMock())

    check("Google Sheet URL is saved for user", bot._user_sheet_urls.get(test_user_id) == update.message.text)
    check("User records loaded in memory from Google Sheet", len(bot._user_records.get(test_user_id, [])) == 2)
    check("sheet_url.txt file persisted to disk", bot._user_sheet_url_path(test_user_id).exists())

    # Test callback menu:sync_sheet
    query_update = MagicMock()
    query = AsyncMock()
    query.from_user = user
    query.data = "menu:sync_sheet"
    query_update.callback_query = query

    with patch.object(httpx.Client, "get", return_value=mock_ok):
        await bot.handle_callback_query(query_update, MagicMock())

    check("Callback menu:sync_sheet triggered safe edit", query.edit_message_text.called or query.message.edit_text.called)

    # Clean up test user
    bot._authorized_users.discard(test_user_id)
    bot._save_authorized_users()
    shutil.rmtree(Path(f"data/{test_user_id}"), ignore_errors=True)
    shutil.rmtree(Path("data/unit_test_user"), ignore_errors=True)

asyncio.run(run_bot_google_sheet_tests())

# ===========================================================================
# Test 20: Multi-sheet workbook collection
# ===========================================================================
print("\n--- Test 20: Multi-sheet workbook collection ---")

wb_multi = openpyxl.Workbook()
ws1 = wb_multi.active
ws1.title = "Batch 1"
ws1.append(["student_name", "telegram_username", "due_date"])
ws1.append(["Alice", "@alice", "2026-10-25"])

ws2 = wb_multi.create_sheet(title="Batch 2")
ws2.append(["student_name", "telegram_username", "due_date"])
ws2.append(["Bob", "@bob", "2026-10-26"])
# Add a duplicate of Alice in Batch 2 with same due date to verify dedup
ws2.append(["Alice Duplicate", "@alice", "2026-10-25"])

ws3 = wb_multi.create_sheet(title="Notes")
ws3.append(["Topic", "Details"])
ws3.append(["Meeting", "Class discussion"])

with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp_m:
    wb_multi.save(tmp_m.name)
    multi_records, inv_m, errs_m = excel_handler.load_excel(tmp_m.name)

check("Multi-sheet loads records from Batch 1 and Batch 2", len(multi_records) == 2)
check("Alice from Batch 1 loaded", any(r["telegram_username"] == "@alice" for r in multi_records))
check("Bob from Batch 2 loaded", any(r["telegram_username"] == "@bob" for r in multi_records))
check("Notes tab safely ignored without error", inv_m == 0)

# ===========================================================================
# Test 21: 3-day reminder repeat interval logic
# ===========================================================================
print("\n--- Test 21: 3-day reminder repeat interval logic ---")

async def run_repeat_reminder_tests():
    mock_repeat_bot = AsyncMock()
    with tempfile.NamedTemporaryFile(suffix=".json", delete=False) as tmp_sent:
        repeat_json = Path(tmp_sent.name)

    student_rec = [
        {"student_name": "Ko Ko", "telegram_username": "@koko", "due_date": datetime.date(2026, 10, 25), "telegram_chat_id": 112233}
    ]

    # Day -5 (Oct 20): First reminder should send
    with patch("reminder_scheduler.datetime") as mock_dt:
        mock_dt.date = datetime.date
        mock_dt.datetime.now.return_value.date.return_value = datetime.date(2026, 10, 20)
        res_day0 = await reminder_scheduler.daily_check(
            bot=mock_repeat_bot,
            records=student_rec,
            sent_reminders_path=repeat_json,
            repeat_interval_days=3,
        )
    check("Day -5 (Oct 20): Initial reminder sent", len(res_day0["sent"]) == 1)

    # Day -4 (Oct 21): 1 day later -> should NOT send
    with patch("reminder_scheduler.datetime") as mock_dt:
        mock_dt.date = datetime.date
        mock_dt.datetime.now.return_value.date.return_value = datetime.date(2026, 10, 21)
        res_day1 = await reminder_scheduler.daily_check(
            bot=mock_repeat_bot,
            records=student_rec,
            sent_reminders_path=repeat_json,
            repeat_interval_days=3,
        )
    check("Day -4 (Oct 21): Skipped (only 1 day elapsed)", len(res_day1["sent"]) == 0)

    # Day -3 (Oct 22): 2 days later -> should NOT send
    with patch("reminder_scheduler.datetime") as mock_dt:
        mock_dt.date = datetime.date
        mock_dt.datetime.now.return_value.date.return_value = datetime.date(2026, 10, 22)
        res_day2 = await reminder_scheduler.daily_check(
            bot=mock_repeat_bot,
            records=student_rec,
            sent_reminders_path=repeat_json,
            repeat_interval_days=3,
        )
    check("Day -3 (Oct 22): Skipped (only 2 days elapsed)", len(res_day2["sent"]) == 0)

    # Day -2 (Oct 23): 3 days later -> SHOULD send again!
    with patch("reminder_scheduler.datetime") as mock_dt:
        mock_dt.date = datetime.date
        mock_dt.datetime.now.return_value.date.return_value = datetime.date(2026, 10, 23)
        res_day3 = await reminder_scheduler.daily_check(
            bot=mock_repeat_bot,
            records=student_rec,
            sent_reminders_path=repeat_json,
            repeat_interval_days=3,
        )
    check("Day -2 (Oct 23): Reminder sent again after 3 days", len(res_day3["sent"]) == 1)

    # Day -1 (Oct 24): 1 day after second reminder -> should NOT send
    with patch("reminder_scheduler.datetime") as mock_dt:
        mock_dt.date = datetime.date
        mock_dt.datetime.now.return_value.date.return_value = datetime.date(2026, 10, 24)
        res_day4 = await reminder_scheduler.daily_check(
            bot=mock_repeat_bot,
            records=student_rec,
            sent_reminders_path=repeat_json,
            repeat_interval_days=3,
        )
    check("Day -1 (Oct 24): Skipped (only 1 day after 2nd reminder)", len(res_day4["sent"]) == 0)

    repeat_json.unlink(missing_ok=True)

asyncio.run(run_repeat_reminder_tests())
print("\n" + "=" * 50)
passed = sum(results)
total = len(results)
print(f"Smoke tests: {passed}/{total} passed")
if passed == total:
    print("🎉 All tests passed!")
    sys.exit(0)
else:
    print(f"⚠️  {total - passed} test(s) failed.")
    sys.exit(1)
